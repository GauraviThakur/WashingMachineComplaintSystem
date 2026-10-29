
import csv
import io
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError

from app.repositories.bulk_request_repository import (
    create_bulk_request,
    create_bulk_request_rows,
    finish_bulk_request,
    get_bulk_request,
    get_bulk_request_rows,
    update_bulk_request_row,
)
from app.schemas.complaint import ComplaintCreate
from app.services.complaint_service import create_complaint


REQUIRED_COLUMNS = {
    "serial_number",
    "customer_name",
    "dealer_email",
    "reported_issue",
}


def _normalize_row(
    row: dict[str, Any],
) -> dict[str, Any]:
    """Normalize blank optional values and apply defaults."""

    normalized = {
        str(key).strip().lower(): (
            value.strip() if isinstance(value, str) else value
        )
        for key, value in row.items()
        if key is not None
    }

    for key, value in normalized.items():
        if value == "":
            normalized[key] = None

    if normalized.get("issue_type_id") is None:
        normalized["issue_type_id"] = None
    else:
        try:
            normalized["issue_type_id"] = int(
                normalized["issue_type_id"]
            )
        except (ValueError, TypeError):
            pass

    if not normalized.get("registered_by"):
        normalized["registered_by"] = "system"

    return normalized


def _read_csv(content: bytes) -> list[dict[str, Any]]:
    """Read CSV content and return rows as dictionaries."""

    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError(
            "CSV must use UTF-8 encoding."
        ) from exc

    reader = csv.DictReader(io.StringIO(text))

    if not reader.fieldnames:
        raise ValueError("The uploaded CSV file is empty.")

    headers = {
        str(header).strip().lower()
        for header in reader.fieldnames
        if header is not None
    }

    missing = REQUIRED_COLUMNS - headers
    if missing:
        raise ValueError(
            "Missing required columns: "
            + ", ".join(sorted(missing))
        )

    rows = []

    for row_number, row in enumerate(reader, start=2):
        normalized = _normalize_row(row)

        if not any(value is not None for value in normalized.values()):
            continue

        rows.append(
            {
                "row_number": row_number,
                "row_data": normalized,
            }
        )

    return rows


def _read_excel(content: bytes) -> list[dict[str, Any]]:
    """Read the first worksheet of an Excel workbook."""

    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ValueError(
            "Unable to read the Excel file. "
            "Please upload a valid .xlsx workbook."
        ) from exc

    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        headers_row = next(iterator, None)

        if not headers_row:
            raise ValueError("The uploaded Excel file is empty.")

        headers = [
            str(value).strip().lower() if value is not None else ""
            for value in headers_row
        ]

        missing = REQUIRED_COLUMNS - set(headers)
        if missing:
            raise ValueError(
                "Missing required columns: "
                + ", ".join(sorted(missing))
            )

        nonempty_headers = [header for header in headers if header]
        if len(nonempty_headers) != len(set(nonempty_headers)):
            raise ValueError(
                "The Excel file contains duplicate column names."
            )

        rows = []

        for row_number, values in enumerate(iterator, start=2):
            raw_row = {
                header: value
                for header, value in zip(headers, values)
                if header
            }

            normalized = _normalize_row(raw_row)

            if not any(value is not None for value in normalized.values()):
                continue

            rows.append(
                {
                    "row_number": row_number,
                    "row_data": normalized,
                }
            )

        return rows
    finally:
        workbook.close()


def _read_upload(
    *,
    filename: str,
    content: bytes,
) -> list[dict[str, Any]]:
    """Select the reader based on the uploaded file extension."""

    extension = (
        filename.rsplit(".", 1)[-1].lower()
        if "." in filename
        else ""
    )

    if extension == "csv":
        return _read_csv(content)

    if extension == "xlsx":
        return _read_excel(content)

    raise ValueError(
        "Unsupported file format. Upload a .csv or .xlsx file."
    )


def process_bulk_upload(
    *,
    filename: str,
    content: bytes,
    uploaded_by: str,
) -> dict[str, Any]:
    """Validate uploaded rows and register valid complaints."""

    if not content:
        raise ValueError("The uploaded file is empty.")

    uploaded_rows = _read_upload(
        filename=filename,
        content=content,
    )

    if not uploaded_rows:
        raise ValueError(
            "The uploaded file contains no complaint rows."
        )

    bulk_request = create_bulk_request(
        source_file_name=filename,
        uploaded_by=uploaded_by,
        total_rows=len(uploaded_rows),
    )

    bulk_request_id = bulk_request["bulk_request_id"]

    create_bulk_request_rows(
        bulk_request_id=bulk_request_id,
        rows=uploaded_rows,
    )

    successful_count = 0
    failed_count = 0
    results = []

    for uploaded_row in uploaded_rows:
        row_number = uploaded_row["row_number"]
        row_data = uploaded_row["row_data"]

        try:
            complaint_data = ComplaintCreate.model_validate(
                row_data
            )
        except ValidationError as exc:
            errors = [
                (
                    f"{'.'.join(str(part) for part in error['loc'])}: "
                    f"{error['msg']}"
                )
                for error in exc.errors()
            ]

            update_bulk_request_row(
                bulk_request_id=bulk_request_id,
                row_number=row_number,
                validation_status="INVALID",
                validation_errors=errors,
            )

            failed_count += 1
            results.append(
                {
                    "row_number": row_number,
                    "status": "INVALID",
                    "errors": errors,
                }
            )
            continue

        complaint_result = create_complaint(
            {
                **complaint_data.model_dump(),
                "bulk_request_id": bulk_request_id,
            }
        )

        error = complaint_result.get("error")

        if error:
            if error == "machine_not_found":
                errors = [
                    "Machine serial number was not found in manufacturing data."
                ]
            elif error == "active_complaint_exists":
                existing = complaint_result.get(
                    "existing_complaint",
                    {},
                )
                errors = [
                    "An active complaint already exists: "
                    f"{existing.get('complaint_number', 'unknown')}."
                ]
            elif error == "invalid_issue_type":
                errors = [
                    "The selected issue type is invalid or inactive."
                ]
            else:
                errors = ["The complaint could not be registered."]

            update_bulk_request_row(
                bulk_request_id=bulk_request_id,
                row_number=row_number,
                validation_status="INVALID",
                validation_errors=errors,
            )

            failed_count += 1
            results.append(
                {
                    "row_number": row_number,
                    "status": "INVALID",
                    "errors": errors,
                }
            )
            continue

        update_bulk_request_row(
            bulk_request_id=bulk_request_id,
            row_number=row_number,
            validation_status="PROCESSED",
            complaint_id=complaint_result["complaint_id"],
        )

        successful_count += 1
        results.append(
            {
                "row_number": row_number,
                "status": "PROCESSED",
                "complaint_number": complaint_result[
                    "complaint_number"
                ],
                "complaint_id": complaint_result["complaint_id"],
            }
        )

    final_request = finish_bulk_request(
        bulk_request_id=bulk_request_id,
        valid_rows=successful_count,
        invalid_rows=failed_count,
    )

    return {
        "message": "Bulk upload processing completed.",
        "bulk_request": final_request,
        "row_results": results,
    }


def get_bulk_request_summary(
    bulk_request_id: int,
) -> dict[str, Any] | None:
    return get_bulk_request(bulk_request_id)


def get_bulk_request_details(
    bulk_request_id: int,
) -> dict[str, Any] | None:
    request = get_bulk_request(bulk_request_id)

    if request is None:
        return None

    rows = get_bulk_request_rows(bulk_request_id)

    return {
        "bulk_request": request,
        "rows": rows or [],
    }
