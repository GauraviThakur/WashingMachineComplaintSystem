
from datetime import date, datetime, time
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.bulk_request_service import get_bulk_request_details


def _excel_value(value: Any) -> Any:
    """Convert database values into values supported by Excel."""

    if isinstance(value, datetime):
        # Excel does not support timezone-aware datetimes.
        if value.tzinfo is not None:
            value = value.replace(tzinfo=None)
        return value

    if isinstance(value, time):
        if value.tzinfo is not None:
            value = value.replace(tzinfo=None)
        return value

    if isinstance(value, date):
        return value

    if isinstance(value, (list, tuple)):
        return "; ".join(str(item) for item in value)

    if isinstance(value, dict):
        return "; ".join(
            f"{key}: {item}" for key, item in value.items()
        )

    return value


def _display_value(value: Any) -> str:
    """Convert values such as JSON errors into readable Excel text."""

    if value is None:
        return ""

    if isinstance(value, (list, tuple)):
        return "; ".join(str(item) for item in value)

    if isinstance(value, dict):
        return "; ".join(
            f"{key}: {item}" for key, item in value.items()
        )

    return str(value)


def _format_worksheet(worksheet) -> None:
    """Apply readable formatting to a worksheet."""

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 32

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        maximum_length = 0

        for cell in column_cells:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            if cell.value is not None:
                maximum_length = max(
                    maximum_length,
                    max(
                        (
                            len(line)
                            for line in str(cell.value).splitlines()
                        ),
                        default=0,
                    ),
                )

        worksheet.column_dimensions[column_letter].width = min(
            max(maximum_length + 2, 14),
            45,
        )


def generate_bulk_request_report(
    bulk_request_id: int,
) -> BytesIO | None:
    """Generate a two-sheet Excel report from stored database records."""

    details = get_bulk_request_details(bulk_request_id)

    if details is None:
        return None

    request = details["bulk_request"]
    rows = details["rows"]

    workbook = Workbook()

    # Worksheet 1: Summary
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Field", "Value"])

    summary_fields = [
        ("Bulk Request ID", request.get("bulk_request_id")),
        ("Source Filename", request.get("source_file_name")),
        ("Uploaded By", request.get("uploaded_by")),
        ("Upload Date", request.get("uploaded_at")),
        ("Completion Date", request.get("completed_at")),
        ("Total Rows", request.get("total_rows", 0)),
        ("Valid Rows", request.get("valid_rows", 0)),
        ("Invalid Rows", request.get("invalid_rows", 0)),
        ("Request Status", request.get("request_status")),
    ]

    for field, value in summary_fields:
        summary_sheet.append([field, _excel_value(value)])

    # Worksheet 2: Row Details
    details_sheet = workbook.create_sheet("Row Details")

    headers = [
        "Source Row Number",
        "Machine Serial Number",
        "Customer Name",
        "Dealer Email",
        "Reported Issue",
        "Issue Type",
        "Validation Result",
        "Validation Errors",
        "Complaint Number",
        "Complaint Status",
        "Priority",
        "Processing Result",
    ]

    details_sheet.append(headers)

    for row in rows:
        row_data = row.get("row_data") or {}

        if not isinstance(row_data, dict):
            row_data = {}

        validation_status = row.get("validation_status") or "PENDING"
        complaint_number = row.get("complaint_number")

        validation_errors = _display_value(
            row.get("validation_errors")
        )

        if validation_status == "PROCESSED" and complaint_number:
            processing_result = "Complaint registered successfully"
        elif validation_status == "INVALID":
            processing_result = (
                "Row rejected during validation or registration"
            )
        else:
            processing_result = (
                "Not processed or processing incomplete"
            )

        details_sheet.append(
            [
                _excel_value(row.get("row_number")),
                _display_value(row_data.get("serial_number")),
                _display_value(row_data.get("customer_name")),
                _display_value(row_data.get("dealer_email")),
                _display_value(
                    row_data.get("reported_issue")
                    or row.get("registered_reported_issue")
                ),
                _display_value(
                    row.get("issue_type_name")
                    or row_data.get("issue_type_id")
                ),
                _display_value(validation_status),
                validation_errors,
                _display_value(complaint_number),
                _display_value(row.get("complaint_status")),
                _display_value(row.get("priority")),
                processing_result,
            ]
        )

    # Format both worksheets.
    _format_worksheet(summary_sheet)
    _format_worksheet(details_sheet)

    # Format summary timestamps for readability.
    for row_number in range(2, summary_sheet.max_row + 1):
        field_name = summary_sheet.cell(
            row=row_number,
            column=1,
        ).value

        if field_name in ("Upload Date", "Completion Date"):
            cell = summary_sheet.cell(
                row=row_number,
                column=2,
            )
            if cell.value is not None:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)

    return output
